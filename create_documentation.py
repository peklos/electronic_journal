#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_documentation():
    # Создаем документ
    doc = Document()

    # Настройка стилей
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Заголовок
    title = doc.add_heading('Электронный журнал классного руководителя БППК', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    subtitle = doc.add_paragraph('Техническая документация')
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.bold = True

    doc.add_page_break()

    # Содержание
    doc.add_heading('Содержание', 1)
    toc = doc.add_paragraph()
    toc.add_run('1. Введение\n')
    toc.add_run('2. Окно входа (LoginWindow)\n')
    toc.add_run('3. Окно регистрации (RegisterWindow)\n')
    toc.add_run('4. Панель учителя (TeacherWindow)\n')
    toc.add_run('5. Панель ученика (StudentWindow)\n')
    toc.add_run('6. Диалоговые окна\n')

    doc.add_page_break()

    # 1. Введение
    doc.add_heading('1. Введение', 1)
    intro = doc.add_paragraph()
    intro.add_run('Электронный журнал классного руководителя БППК - это WPF приложение на .NET 8.0, '
                  'предназначенное для управления классом, учета оценок и посещаемости учеников.\n\n')
    intro.add_run('Технологии:\n')
    intro.add_run('• Frontend: WPF (Windows Presentation Foundation)\n')
    intro.add_run('• Backend: .NET 8.0\n')
    intro.add_run('• База данных: SQLite\n\n')
    intro.add_run('Основные функции:\n')
    intro.add_run('• Управление учениками\n')
    intro.add_run('• Выставление оценок\n')
    intro.add_run('• Учет посещаемости\n')
    intro.add_run('• Просмотр статистики\n')

    doc.add_page_break()

    # 2. Окно входа
    doc.add_heading('2. Окно входа (LoginWindow)', 1)

    doc.add_heading('2.1. Описание', 2)
    p = doc.add_paragraph('Окно входа предоставляет интерфейс для авторизации пользователей в системе. '
                          'Поддерживает два типа ролей: Учитель и Ученик.')

    doc.add_heading('2.2. Интерфейс', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ОКНА ВХОДА]\n').bold = True
    p.add_run('\n')
    p.add_run('Компоненты интерфейса:\n')
    p.add_run('• Поле ввода "Имя пользователя"\n')
    p.add_run('• Поле ввода "Пароль" (скрытое)\n')
    p.add_run('• Кнопка "Войти"\n')
    p.add_run('• Ссылка "Зарегистрироваться"\n')
    p.add_run('• Блок с тестовыми данными для входа\n')

    doc.add_heading('2.3. Функциональность', 2)
    p = doc.add_paragraph()
    p.add_run('• Проверка заполненности полей\n')
    p.add_run('• Аутентификация через DatabaseService\n')
    p.add_run('• Перенаправление на соответствующее окно в зависимости от роли\n')
    p.add_run('• Отображение ошибок авторизации\n')

    doc.add_heading('2.4. Код XAML (интерфейс)', 2)
    xaml_code = '''<Window x:Class="ElectronicJournal.Views.LoginWindow"
        xmlns="http://schemas.microsoft.com/winfx/2006/xaml/presentation"
        xmlns:x="http://schemas.microsoft.com/winfx/2006/xaml"
        Title="Электронный журнал БППК - Вход"
        Height="650" Width="500"
        WindowStartupLocation="CenterScreen"
        Background="{StaticResource BackgroundColor}">
    <Grid>
        <!-- Шапка -->
        <Border Background="{StaticResource PrimaryColor}" Padding="20">
            <StackPanel>
                <TextBlock Text="Электронный журнал"
                          FontSize="24" FontWeight="Bold"/>
                <TextBlock Text="БППК - Система управления классом"
                          FontSize="12"/>
            </StackPanel>
        </Border>

        <!-- Форма входа -->
        <Border Background="{StaticResource SurfaceColor}">
            <StackPanel>
                <TextBox x:Name="UsernameTextBox"/>
                <PasswordBox x:Name="PasswordBox"/>
                <TextBlock x:Name="ErrorTextBlock"/>
                <Button Content="Войти" Click="LoginButton_Click"/>
                <Hyperlink Click="RegisterLink_Click">
                    Зарегистрироваться
                </Hyperlink>
            </StackPanel>
        </Border>
    </Grid>
</Window>'''
    p = doc.add_paragraph(xaml_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_heading('2.5. Код C# (логика)', 2)
    cs_code = '''using System.Windows;
using ElectronicJournal.Services;

namespace ElectronicJournal.Views
{
    public partial class LoginWindow : Window
    {
        private readonly DatabaseService _databaseService;

        public LoginWindow()
        {
            InitializeComponent();
            _databaseService = new DatabaseService();
        }

        private void LoginButton_Click(object sender, RoutedEventArgs e)
        {
            string username = UsernameTextBox.Text.Trim();
            string password = PasswordBox.Password;

            if (string.IsNullOrEmpty(username) || string.IsNullOrEmpty(password))
            {
                ErrorTextBlock.Text = "Пожалуйста, заполните все поля";
                return;
            }

            var user = _databaseService.AuthenticateUser(username, password);

            if (user != null)
            {
                if (user.Role == "Учитель")
                {
                    var teacherWindow = new TeacherWindow(user);
                    teacherWindow.Show();
                    this.Close();
                }
                else if (user.Role == "Ученик")
                {
                    var studentWindow = new StudentWindow(user);
                    studentWindow.Show();
                    this.Close();
                }
            }
            else
            {
                ErrorTextBlock.Text = "Неверное имя пользователя или пароль";
            }
        }

        private void RegisterLink_Click(object sender, RoutedEventArgs e)
        {
            var registerWindow = new RegisterWindow();
            registerWindow.ShowDialog();
        }
    }
}'''
    p = doc.add_paragraph(cs_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    # 3. Окно регистрации
    doc.add_heading('3. Окно регистрации (RegisterWindow)', 1)

    doc.add_heading('3.1. Описание', 2)
    p = doc.add_paragraph('Окно регистрации позволяет создавать новые учетные записи пользователей. '
                          'Поддерживает регистрацию учителей и учеников.')

    doc.add_heading('3.2. Интерфейс', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ОКНА РЕГИСТРАЦИИ]\n').bold = True
    p.add_run('\n')
    p.add_run('Компоненты интерфейса:\n')
    p.add_run('• Поле "Имя пользователя"\n')
    p.add_run('• Поле "Пароль"\n')
    p.add_run('• Поле "Подтверждение пароля"\n')
    p.add_run('• Поле "ФИО"\n')
    p.add_run('• Выпадающий список "Роль" (Учитель/Ученик)\n')
    p.add_run('• Выпадающий список "Выберите себя из списка учеников" (только для учеников)\n')
    p.add_run('• Кнопка "Зарегистрироваться"\n')
    p.add_run('• Кнопка "Отмена"\n')

    doc.add_heading('3.3. Функциональность', 2)
    p = doc.add_paragraph()
    p.add_run('• Валидация всех полей\n')
    p.add_run('• Проверка совпадения паролей\n')
    p.add_run('• Минимальная длина имени пользователя (3 символа)\n')
    p.add_run('• Минимальная длина пароля (3 символа)\n')
    p.add_run('• Динамическое отображение списка учеников для роли "Ученик"\n')
    p.add_run('• Регистрация через DatabaseService\n')
    p.add_run('• Проверка уникальности имени пользователя\n')

    doc.add_heading('3.4. Код XAML (интерфейс)', 2)
    xaml_code = '''<Window x:Class="ElectronicJournal.Views.RegisterWindow"
        Title="Регистрация - Электронный журнал БППК"
        Height="700" Width="500"
        WindowStartupLocation="CenterScreen">
    <Grid>
        <Border Background="{StaticResource SurfaceColor}">
            <StackPanel>
                <TextBlock Text="Создание нового аккаунта"
                          FontSize="18" FontWeight="SemiBold"/>

                <TextBox x:Name="UsernameTextBox"/>
                <PasswordBox x:Name="PasswordBox"/>
                <PasswordBox x:Name="ConfirmPasswordBox"/>
                <TextBox x:Name="FullNameTextBox"/>

                <ComboBox x:Name="RoleComboBox"
                         SelectionChanged="RoleComboBox_SelectionChanged">
                    <ComboBoxItem Content="Учитель" IsSelected="True"/>
                    <ComboBoxItem Content="Ученик"/>
                </ComboBox>

                <StackPanel x:Name="StudentSelectionPanel"
                           Visibility="Collapsed">
                    <ComboBox x:Name="StudentComboBox"
                             DisplayMemberPath="FullName"/>
                </StackPanel>

                <TextBlock x:Name="ErrorTextBlock"/>
                <Button Content="Зарегистрироваться"
                       Click="RegisterButton_Click"/>
                <Button Content="Отмена" Click="CancelButton_Click"/>
            </StackPanel>
        </Border>
    </Grid>
</Window>'''
    p = doc.add_paragraph(xaml_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_heading('3.5. Код C# (логика)', 2)
    cs_code = '''using System.Windows;
using System.Windows.Controls;
using ElectronicJournal.Services;

namespace ElectronicJournal.Views
{
    public partial class RegisterWindow : Window
    {
        private readonly DatabaseService _databaseService;

        public RegisterWindow()
        {
            InitializeComponent();
            _databaseService = new DatabaseService();
            LoadStudents();
        }

        private void LoadStudents()
        {
            var students = _databaseService.GetAllStudents();
            StudentComboBox.ItemsSource = students;
        }

        private void RoleComboBox_SelectionChanged(object sender, SelectionChangedEventArgs e)
        {
            if (StudentSelectionPanel == null) return;

            if (RoleComboBox.SelectedIndex == 1) // Ученик
            {
                StudentSelectionPanel.Visibility = Visibility.Visible;
            }
            else
            {
                StudentSelectionPanel.Visibility = Visibility.Collapsed;
            }
        }

        private void RegisterButton_Click(object sender, RoutedEventArgs e)
        {
            string username = UsernameTextBox.Text.Trim();
            string password = PasswordBox.Password;
            string confirmPassword = ConfirmPasswordBox.Password;
            string fullName = FullNameTextBox.Text.Trim();
            string role = ((ComboBoxItem)RoleComboBox.SelectedItem).Content.ToString();

            // Валидация полей
            if (string.IsNullOrEmpty(username) || string.IsNullOrEmpty(password) ||
                string.IsNullOrEmpty(fullName))
            {
                ErrorTextBlock.Text = "Пожалуйста, заполните все поля";
                return;
            }

            if (password != confirmPassword)
            {
                ErrorTextBlock.Text = "Пароли не совпадают";
                return;
            }

            if (username.Length < 3)
            {
                ErrorTextBlock.Text = "Имя пользователя должно быть не менее 3 символов";
                return;
            }

            if (password.Length < 3)
            {
                ErrorTextBlock.Text = "Пароль должен быть не менее 3 символов";
                return;
            }

            int? studentId = null;
            if (role == "Ученик")
            {
                if (StudentComboBox.SelectedItem == null)
                {
                    ErrorTextBlock.Text = "Пожалуйста, выберите себя из списка учеников";
                    return;
                }
                studentId = ((Models.Student)StudentComboBox.SelectedItem).Id;
            }

            bool success = _databaseService.RegisterUser(username, password, fullName, role, studentId);

            if (success)
            {
                MessageBox.Show("Регистрация прошла успешно!");
                this.Close();
            }
            else
            {
                ErrorTextBlock.Text = "Пользователь с таким именем уже существует";
            }
        }
    }
}'''
    p = doc.add_paragraph(cs_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    # 4. Панель учителя
    doc.add_heading('4. Панель учителя (TeacherWindow)', 1)

    doc.add_heading('4.1. Описание', 2)
    p = doc.add_paragraph('Панель учителя - основное окно управления классом. Содержит три вкладки: '
                          'Ученики, Оценки и Посещаемость.')

    doc.add_heading('4.2. Интерфейс', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ПАНЕЛИ УЧИТЕЛЯ - ОБЩИЙ ВИД]\n').bold = True

    doc.add_heading('4.3. Вкладка "Ученики"', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ВКЛАДКИ УЧЕНИКИ]\n').bold = True
    p.add_run('\n')
    p.add_run('Функции:\n')
    p.add_run('• Просмотр списка всех учеников\n')
    p.add_run('• Добавление нового ученика\n')
    p.add_run('• Редактирование данных ученика\n')
    p.add_run('• Удаление ученика\n')
    p.add_run('• Обновление списка\n')
    p.add_run('\n')
    p.add_run('Отображаемые данные:\n')
    p.add_run('• ID\n')
    p.add_run('• ФИО\n')
    p.add_run('• Класс\n')
    p.add_run('• Дата рождения\n')
    p.add_run('• Телефон родителя\n')
    p.add_run('• Адрес\n')

    doc.add_heading('4.4. Вкладка "Оценки"', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ВКЛАДКИ ОЦЕНКИ]\n').bold = True
    p.add_run('\n')
    p.add_run('Функции:\n')
    p.add_run('• Просмотр всех оценок\n')
    p.add_run('• Добавление новой оценки\n')
    p.add_run('• Удаление оценки\n')
    p.add_run('• Фильтрация по ученику\n')
    p.add_run('• Обновление списка\n')
    p.add_run('\n')
    p.add_run('Отображаемые данные:\n')
    p.add_run('• ID\n')
    p.add_run('• Ученик\n')
    p.add_run('• Предмет\n')
    p.add_run('• Оценка (с цветовым кодированием)\n')
    p.add_run('• Дата\n')
    p.add_run('• Тема урока\n')
    p.add_run('• Примечание\n')
    p.add_run('\n')
    p.add_run('Цветовое кодирование оценок:\n')
    p.add_run('• 5 - Зеленый (#27AE60)\n')
    p.add_run('• 4 - Синий (#3498DB)\n')
    p.add_run('• 3 - Оранжевый (#F39C12)\n')
    p.add_run('• 2 - Красный (#E74C3C)\n')

    doc.add_heading('4.5. Вкладка "Посещаемость"', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ВКЛАДКИ ПОСЕЩАЕМОСТЬ]\n').bold = True
    p.add_run('\n')
    p.add_run('Функции:\n')
    p.add_run('• Просмотр записей посещаемости\n')
    p.add_run('• Отметка посещаемости\n')
    p.add_run('• Удаление записи\n')
    p.add_run('• Фильтрация по ученику\n')
    p.add_run('• Обновление списка\n')
    p.add_run('\n')
    p.add_run('Отображаемые данные:\n')
    p.add_run('• ID\n')
    p.add_run('• Ученик\n')
    p.add_run('• Дата\n')
    p.add_run('• Статус (с цветовым кодированием)\n')
    p.add_run('• Причина\n')
    p.add_run('\n')
    p.add_run('Статусы посещаемости:\n')
    p.add_run('• Присутствовал - Зеленый (#27AE60)\n')
    p.add_run('• Отсутствовал - Красный (#E74C3C)\n')
    p.add_run('• По уважительной причине - Оранжевый (#F39C12)\n')

    doc.add_heading('4.6. Код XAML (фрагмент)', 2)
    xaml_code = '''<Window x:Class="ElectronicJournal.Views.TeacherWindow"
        Title="Электронный журнал - Панель учителя"
        Height="700" Width="1200"
        WindowState="Maximized">
    <Grid>
        <!-- Шапка -->
        <Border Background="{StaticResource PrimaryColor}">
            <Grid>
                <StackPanel>
                    <TextBlock x:Name="WelcomeTextBlock"
                              FontSize="20" FontWeight="Bold"/>
                    <TextBlock Text="Панель управления классного руководителя"
                              FontSize="12"/>
                </StackPanel>
                <Button Content="Выход" Click="LogoutButton_Click"/>
            </Grid>
        </Border>

        <!-- Вкладки -->
        <TabControl>
            <!-- Вкладка: Ученики -->
            <TabItem Header="Ученики">
                <Grid>
                    <StackPanel>
                        <Button Content="Добавить ученика"
                               Click="AddStudentButton_Click"/>
                        <Button Content="Редактировать"
                               Click="EditStudentButton_Click"/>
                        <Button Content="Удалить"
                               Click="DeleteStudentButton_Click"/>
                    </StackPanel>
                    <DataGrid x:Name="StudentsDataGrid"/>
                </Grid>
            </TabItem>

            <!-- Вкладка: Оценки -->
            <TabItem Header="Оценки">
                <Grid>
                    <StackPanel>
                        <Button Content="Добавить оценку"
                               Click="AddGradeButton_Click"/>
                        <ComboBox x:Name="GradeFilterComboBox"
                                 SelectionChanged="GradeFilterComboBox_SelectionChanged"/>
                    </StackPanel>
                    <DataGrid x:Name="GradesDataGrid"/>
                </Grid>
            </TabItem>

            <!-- Вкладка: Посещаемость -->
            <TabItem Header="Посещаемость">
                <Grid>
                    <StackPanel>
                        <Button Content="Отметить посещаемость"
                               Click="AddAttendanceButton_Click"/>
                        <ComboBox x:Name="AttendanceFilterComboBox"
                                 SelectionChanged="AttendanceFilterComboBox_SelectionChanged"/>
                    </StackPanel>
                    <DataGrid x:Name="AttendanceDataGrid"/>
                </Grid>
            </TabItem>
        </TabControl>
    </Grid>
</Window>'''
    p = doc.add_paragraph(xaml_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    doc.add_heading('4.7. Код C# (основные методы)', 2)
    cs_code = '''using System.Windows;
using ElectronicJournal.Models;
using ElectronicJournal.Services;

namespace ElectronicJournal.Views
{
    public partial class TeacherWindow : Window
    {
        private readonly DatabaseService _databaseService;
        private readonly User _currentUser;

        public TeacherWindow(User user)
        {
            InitializeComponent();
            _databaseService = new DatabaseService();
            _currentUser = user;
            WelcomeTextBlock.Text = $"Добро пожаловать, {_currentUser.FullName}!";
            LoadData();
        }

        private void LoadData()
        {
            LoadStudents();
            LoadGrades();
            LoadAttendance();
        }

        private void LoadStudents()
        {
            var students = _databaseService.GetAllStudents();
            StudentsDataGrid.ItemsSource = students;
            GradeFilterComboBox.ItemsSource = students;
            AttendanceFilterComboBox.ItemsSource = students;
            StudentsCountTextBlock.Text = $"Всего учеников: {students.Count}";
        }

        private void AddStudentButton_Click(object sender, RoutedEventArgs e)
        {
            var dialog = new AddEditStudentDialog();
            if (dialog.ShowDialog() == true)
            {
                LoadStudents();
            }
        }

        private void EditStudentButton_Click(object sender, RoutedEventArgs e)
        {
            if (StudentsDataGrid.SelectedItem is Student selectedStudent)
            {
                var dialog = new AddEditStudentDialog(selectedStudent);
                if (dialog.ShowDialog() == true)
                {
                    LoadStudents();
                }
            }
        }

        private void DeleteStudentButton_Click(object sender, RoutedEventArgs e)
        {
            if (StudentsDataGrid.SelectedItem is Student selectedStudent)
            {
                var result = MessageBox.Show(
                    $"Вы уверены, что хотите удалить ученика {selectedStudent.FullName}?",
                    "Подтверждение",
                    MessageBoxButton.YesNo,
                    MessageBoxImage.Question);

                if (result == MessageBoxResult.Yes)
                {
                    _databaseService.DeleteStudent(selectedStudent.Id);
                    LoadStudents();
                }
            }
        }

        // Аналогичные методы для оценок и посещаемости
    }
}'''
    p = doc.add_paragraph(cs_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    # 5. Панель ученика
    doc.add_heading('5. Панель ученика (StudentWindow)', 1)

    doc.add_heading('5.1. Описание', 2)
    p = doc.add_paragraph('Панель ученика предоставляет доступ к личной информации, оценкам и посещаемости. '
                          'Содержит три вкладки для просмотра различных данных.')

    doc.add_heading('5.2. Интерфейс', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ПАНЕЛИ УЧЕНИКА - ОБЩИЙ ВИД]\n').bold = True

    doc.add_heading('5.3. Вкладка "Личная информация"', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ВКЛАДКИ ЛИЧНАЯ ИНФОРМАЦИЯ]\n').bold = True
    p.add_run('\n')
    p.add_run('Отображаемая информация:\n')
    p.add_run('• ФИО\n')
    p.add_run('• Класс\n')
    p.add_run('• Дата рождения\n')
    p.add_run('• Телефон родителя\n')
    p.add_run('• Адрес\n')
    p.add_run('• Примечания\n')
    p.add_run('\n')
    p.add_run('Статистика:\n')
    p.add_run('• Средний балл\n')
    p.add_run('• Всего оценок\n')

    doc.add_heading('5.4. Вкладка "Мои оценки"', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ВКЛАДКИ МОИ ОЦЕНКИ]\n').bold = True
    p.add_run('\n')
    p.add_run('Функции:\n')
    p.add_run('• Просмотр всех оценок ученика\n')
    p.add_run('• Обновление списка\n')
    p.add_run('\n')
    p.add_run('Отображаемые данные:\n')
    p.add_run('• Предмет\n')
    p.add_run('• Оценка (с цветовым кодированием)\n')
    p.add_run('• Дата\n')
    p.add_run('• Тема урока\n')
    p.add_run('• Примечание\n')

    doc.add_heading('5.5. Вкладка "Моя посещаемость"', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ВКЛАДКИ МОЯ ПОСЕЩАЕМОСТЬ]\n').bold = True
    p.add_run('\n')
    p.add_run('Функции:\n')
    p.add_run('• Просмотр истории посещаемости\n')
    p.add_run('• Обновление списка\n')
    p.add_run('• Статистика посещаемости\n')
    p.add_run('\n')
    p.add_run('Отображаемые данные:\n')
    p.add_run('• Дата\n')
    p.add_run('• Статус (с цветовым кодированием)\n')
    p.add_run('• Причина\n')

    doc.add_heading('5.6. Код XAML (фрагмент)', 2)
    xaml_code = '''<Window x:Class="ElectronicJournal.Views.StudentWindow"
        Title="Электронный журнал - Панель ученика"
        Height="700" Width="1000"
        WindowState="Maximized">
    <Grid>
        <!-- Шапка -->
        <Border Background="{StaticResource PrimaryColor}">
            <Grid>
                <StackPanel>
                    <TextBlock x:Name="WelcomeTextBlock"
                              FontSize="20" FontWeight="Bold"/>
                    <TextBlock Text="Личный кабинет ученика"
                              FontSize="12"/>
                </StackPanel>
                <Button Content="Выход" Click="LogoutButton_Click"/>
            </Grid>
        </Border>

        <TabControl>
            <!-- Вкладка: Личная информация -->
            <TabItem Header="Личная информация">
                <StackPanel>
                    <TextBlock x:Name="FullNameTextBlock"/>
                    <TextBlock x:Name="ClassTextBlock"/>
                    <TextBlock x:Name="BirthDateTextBlock"/>

                    <!-- Статистика -->
                    <Border>
                        <StackPanel>
                            <TextBlock x:Name="AverageGradeTextBlock"
                                      FontSize="32" FontWeight="Bold"/>
                            <TextBlock x:Name="TotalGradesTextBlock"
                                      FontSize="32" FontWeight="Bold"/>
                        </StackPanel>
                    </Border>
                </StackPanel>
            </TabItem>

            <!-- Вкладка: Мои оценки -->
            <TabItem Header="Мои оценки">
                <Grid>
                    <Button Content="Обновить"
                           Click="RefreshGradesButton_Click"/>
                    <DataGrid x:Name="GradesDataGrid"/>
                </Grid>
            </TabItem>

            <!-- Вкладка: Моя посещаемость -->
            <TabItem Header="Моя посещаемость">
                <Grid>
                    <Button Content="Обновить"
                           Click="RefreshAttendanceButton_Click"/>
                    <DataGrid x:Name="AttendanceDataGrid"/>
                </Grid>
            </TabItem>
        </TabControl>
    </Grid>
</Window>'''
    p = doc.add_paragraph(xaml_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    doc.add_heading('5.7. Код C# (основные методы)', 2)
    cs_code = '''using System.Linq;
using System.Windows;
using ElectronicJournal.Models;
using ElectronicJournal.Services;

namespace ElectronicJournal.Views
{
    public partial class StudentWindow : Window
    {
        private readonly DatabaseService _databaseService;
        private readonly User _currentUser;
        private Student? _studentInfo;

        public StudentWindow(User user)
        {
            InitializeComponent();
            _databaseService = new DatabaseService();
            _currentUser = user;
            LoadStudentInfo();
            LoadData();
        }

        private void LoadStudentInfo()
        {
            if (_currentUser.StudentId.HasValue)
            {
                _studentInfo = _databaseService.GetStudentById(_currentUser.StudentId.Value);

                if (_studentInfo != null)
                {
                    WelcomeTextBlock.Text = $"Добро пожаловать, {_studentInfo.FullName}!";
                    FullNameTextBlock.Text = _studentInfo.FullName;
                    ClassTextBlock.Text = _studentInfo.Class;
                    BirthDateTextBlock.Text = _studentInfo.BirthDate.ToString("dd.MM.yyyy");
                    ParentPhoneTextBlock.Text = _studentInfo.ParentPhone;
                    AddressTextBlock.Text = _studentInfo.Address;
                }
            }
        }

        private void LoadData()
        {
            LoadGrades();
            LoadAttendance();
            UpdateStatistics();
        }

        private void LoadGrades()
        {
            if (_studentInfo != null)
            {
                var grades = _databaseService.GetGradesByStudent(_studentInfo.Id);
                GradesDataGrid.ItemsSource = grades;
                GradesInfoTextBlock.Text = $"Всего оценок: {grades.Count}";
            }
        }

        private void UpdateStatistics()
        {
            if (_studentInfo != null)
            {
                var grades = _databaseService.GetGradesByStudent(_studentInfo.Id);
                TotalGradesTextBlock.Text = grades.Count.ToString();

                if (grades.Count > 0)
                {
                    double average = grades.Average(g => g.GradeValue);
                    AverageGradeTextBlock.Text = average.ToString("F2");
                }
                else
                {
                    AverageGradeTextBlock.Text = "—";
                }
            }
        }
    }
}'''
    p = doc.add_paragraph(cs_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    # 6. Диалоговые окна
    doc.add_heading('6. Диалоговые окна', 1)

    doc.add_heading('6.1. Добавление/Редактирование ученика (AddEditStudentDialog)', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ДИАЛОГА ДОБАВЛЕНИЯ УЧЕНИКА]\n').bold = True
    p.add_run('\n')
    p.add_run('Поля для ввода:\n')
    p.add_run('• ФИО\n')
    p.add_run('• Класс (например: 9-А)\n')
    p.add_run('• Дата рождения (DatePicker)\n')
    p.add_run('• Телефон родителя\n')
    p.add_run('• Адрес\n')
    p.add_run('• Примечания (многострочное поле)\n')
    p.add_run('\n')
    p.add_run('Кнопки:\n')
    p.add_run('• Сохранить\n')
    p.add_run('• Отмена\n')

    doc.add_heading('6.2. Добавление оценки (AddGradeDialog)', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ДИАЛОГА ДОБАВЛЕНИЯ ОЦЕНКИ]\n').bold = True
    p.add_run('\n')
    p.add_run('Поля для ввода:\n')
    p.add_run('• Ученик (выпадающий список)\n')
    p.add_run('• Предмет (выпадающий список)\n')
    p.add_run('• Оценка (2, 3, 4, 5)\n')
    p.add_run('• Дата (DatePicker)\n')
    p.add_run('• Тема урока\n')
    p.add_run('• Примечание (многострочное поле)\n')
    p.add_run('\n')
    p.add_run('Кнопки:\n')
    p.add_run('• Сохранить\n')
    p.add_run('• Отмена\n')

    doc.add_heading('6.3. Отметка посещаемости (AddAttendanceDialog)', 2)
    p = doc.add_paragraph()
    p.add_run('[ЗДЕСЬ НУЖЕН СКРИНШОТ ДИАЛОГА ОТМЕТКИ ПОСЕЩАЕМОСТИ]\n').bold = True
    p.add_run('\n')
    p.add_run('Поля для ввода:\n')
    p.add_run('• Ученик (выпадающий список)\n')
    p.add_run('• Дата (DatePicker)\n')
    p.add_run('• Статус посещаемости:\n')
    p.add_run('  - Присутствовал\n')
    p.add_run('  - Отсутствовал\n')
    p.add_run('  - По уважительной причине\n')
    p.add_run('• Причина отсутствия (многострочное поле)\n')
    p.add_run('\n')
    p.add_run('Кнопки:\n')
    p.add_run('• Сохранить\n')
    p.add_run('• Отмена\n')

    doc.add_heading('6.4. Код XAML диалогов (фрагменты)', 2)
    xaml_code = '''<!-- AddGradeDialog.xaml -->
<Window x:Class="ElectronicJournal.Views.AddGradeDialog"
        Title="Добавить оценку"
        Height="500" Width="500"
        WindowStartupLocation="CenterOwner">
    <Border>
        <StackPanel>
            <TextBlock Text="Добавить оценку"
                      FontSize="20" FontWeight="SemiBold"/>

            <ComboBox x:Name="StudentComboBox"
                     DisplayMemberPath="FullName"/>
            <ComboBox x:Name="SubjectComboBox"
                     DisplayMemberPath="Name"/>
            <ComboBox x:Name="GradeComboBox">
                <ComboBoxItem Content="5 (Отлично)" Tag="5"/>
                <ComboBoxItem Content="4 (Хорошо)" Tag="4"/>
                <ComboBoxItem Content="3 (Удовлетворительно)" Tag="3"/>
                <ComboBoxItem Content="2 (Неудовлетворительно)" Tag="2"/>
            </ComboBox>
            <DatePicker x:Name="DatePicker"/>
            <TextBox x:Name="TopicTextBox"/>
            <TextBox x:Name="NotesTextBox"
                    Height="60" TextWrapping="Wrap"/>

            <Button Content="Сохранить" Click="SaveButton_Click"/>
            <Button Content="Отмена" Click="CancelButton_Click"/>
        </StackPanel>
    </Border>
</Window>'''
    p = doc.add_paragraph(xaml_code)
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    # Заключение
    doc.add_heading('7. Заключение', 1)
    p = doc.add_paragraph()
    p.add_run('Данная документация описывает все основные окна и функциональность приложения '
              '"Электронный журнал классного руководителя БППК".\n\n')
    p.add_run('Основные возможности:\n')
    p.add_run('• Управление учениками и их данными\n')
    p.add_run('• Выставление и просмотр оценок\n')
    p.add_run('• Учет посещаемости\n')
    p.add_run('• Статистика и аналитика\n')
    p.add_run('• Разграничение прав доступа (Учитель/Ученик)\n')
    p.add_run('\n')
    p.add_run('Технологический стек:\n')
    p.add_run('• WPF для создания современного интерфейса\n')
    p.add_run('• .NET 8.0 для бизнес-логики\n')
    p.add_run('• SQLite для хранения данных\n')
    p.add_run('• Material Design стиль оформления\n')

    # Сохраняем документ
    doc.save('/home/user/electronic_journal/Документация_Электронный_Журнал.docx')
    print("Документация успешно создана!")

if __name__ == '__main__':
    create_documentation()
